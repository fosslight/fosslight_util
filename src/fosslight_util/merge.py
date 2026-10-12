#!/usr/bin/env python
# -*- coding: utf-8 -*-
# Copyright (c) 2020 LG Electronics Inc.
# SPDX-License-Identifier: Apache-2.0

import io
import os
from collections import Counter

import openpyxl

from fosslight_util.constant import COMMENT_DELIMITER, SHEET_NAME_FOR_SCANNER, FOSSLIGHT_SOURCE
from fosslight_util.oss_item import OssItem

SRC_SHEET_NAME = SHEET_NAME_FOR_SCANNER[FOSSLIGHT_SOURCE]
_SRC_COLUMNS = (
    "Source Path",
    "OSS Name",
    "OSS Version",
    "License",
    "Download Location",
    "Homepage",
    "Copyright Text",
    "Exclude",
    "Comment",
)


def _get_source_rows_to_print(source_items: list) -> list:
    source_rows = []
    for source_item in source_items:
        source_rows.extend(source_item.get_print_array())
    return source_rows


def _add_pre_merge_sheet(scan_item, pre_merge_sheet_name: str, header_row: list, pkg_name: str) -> None:
    external_sheets = getattr(scan_item, "external_sheets", {}) or {}
    external_sheets[pre_merge_sheet_name] = [
        header_row,
        *_get_source_rows_to_print(scan_item.file_items.get(pkg_name, [])),
    ]
    scan_item.external_sheets = external_sheets


def _normalize_merge_text(value: str) -> str:
    return value.strip() if value else ""


def _iter_merge_values(values) -> list:
    if not values:
        return []
    if isinstance(values, str):
        return [values]
    return values


def _normalize_merge_licenses(licenses) -> tuple:
    return tuple(sorted([lic.strip() for lic in _iter_merge_values(licenses) if lic and lic.strip()]))


def _normalize_merge_downloads(downloads) -> tuple:
    if not downloads:
        return ()
    if isinstance(downloads, str):
        downloads = [downloads]
    return tuple(sorted([dl.strip() for dl in downloads if dl and dl.strip()]))


def _iter_merge_rows(item):
    if item.oss_items:
        for oss_item in item.oss_items:
            name = _normalize_merge_text(oss_item.name) or _normalize_merge_text(item.oss_name)
            version = _normalize_merge_text(oss_item.version) or _normalize_merge_text(item.oss_version)
            licenses = _normalize_merge_licenses(oss_item.license or item.licenses)
            downloads = _normalize_merge_downloads(oss_item.download_location or item.download_location)
            yield name, version, licenses, downloads
    else:
        yield (
            _normalize_merge_text(item.oss_name),
            _normalize_merge_text(item.oss_version),
            _normalize_merge_licenses(item.licenses),
            _normalize_merge_downloads(item.download_location),
        )


def _get_item_oss_name(item) -> str:
    for name, _, _, _ in _iter_merge_rows(item):
        if name:
            return name
    return ""


def _get_item_oss_version(item) -> str:
    for _, version, _, _ in _iter_merge_rows(item):
        if version:
            return version
    return ""


def _get_merge_licenses(scan_item) -> tuple:
    for _, _, licenses, _ in _iter_merge_rows(scan_item):
        if licenses:
            return licenses
    return ()


def _get_merge_download_locations(scan_item) -> tuple:
    for _, _, _, downloads in _iter_merge_rows(scan_item):
        if downloads:
            return downloads
    return ()


def _get_merge_field_value(scan_items: list, value_getter):
    for scan_item in scan_items:
        value = value_getter(scan_item)
        if value:
            return value
    return ""


def _get_top_merge_values(scan_items: list, value_getter) -> list:
    values = []
    for scan_item in scan_items:
        for value in _iter_merge_values(value_getter(scan_item)):
            normalized_value = _normalize_merge_text(value)
            if normalized_value:
                values.append(normalized_value)
    ranked = sorted(Counter(values).items(), key=lambda entry: (-entry[1], entry[0]))
    return [value for value, _ in ranked[:3]]


def _can_merge_folder(scan_items: list) -> bool:
    if len(scan_items) <= 1:
        return False

    ref_name = None
    ref_version = None
    ref_licenses = None
    ref_downloads = None

    for item in scan_items:
        for name, version, licenses, downloads in _iter_merge_rows(item):
            if name:
                if ref_name is None:
                    ref_name = name
                elif name != ref_name:
                    return False

            if version:
                if ref_version is None:
                    ref_version = version
                elif version != ref_version:
                    return False

            if licenses:
                if ref_licenses is None:
                    ref_licenses = licenses
                elif licenses != ref_licenses:
                    return False

            if downloads:
                if ref_downloads is None:
                    ref_downloads = downloads
                elif downloads != ref_downloads:
                    return False

    return True


def _get_merged_comments(scan_items: list) -> str:
    comments = []
    for item in scan_items:
        val = item.comment
        if val:
            parts = [p.strip() for p in val.split(COMMENT_DELIMITER) if p.strip()]
            for p in parts:
                if p not in comments:
                    comments.append(p)
    if not comments:
        return ""

    return COMMENT_DELIMITER.join(comments)


def _create_merged_item(scan_items: list, merge_path: str):
    merged_path = f"{merge_path} ({len(scan_items)})"
    merged_item = type(scan_items[0])(merged_path)
    merged_item.oss_name = _get_merge_field_value(scan_items, lambda item: _normalize_merge_text(_get_item_oss_name(item)))
    merged_item.oss_version = _get_merge_field_value(scan_items, lambda item: _normalize_merge_text(_get_item_oss_version(item)))
    merged_licenses = _get_merge_field_value(scan_items, _get_merge_licenses)
    merged_item.licenses = list(merged_licenses) if merged_licenses else []
    merged_downloads = _get_merge_field_value(scan_items, _get_merge_download_locations)
    merged_item.download_location = list(merged_downloads) if merged_downloads else []
    merged_copyrights = _get_top_merge_values(scan_items, lambda item: item.copyright)
    merged_item.copyright = merged_copyrights if merged_copyrights else []
    merged_comments = _get_merged_comments(scan_items)
    if merged_comments:
        merged_item.comment = merged_comments
    merged_item.set_oss_item()
    return merged_item


def merge_results_by_folder(scan_result: list) -> list:
    """
    Merge output rows within the same folder when OSS name, OSS version, license,
    and download location are compatible.

    A field is compatible when all non-empty values across rows are identical
    (empty values are ignored). Each oss_item is treated as its own row. All eligible
    rows in the folder must be compatible together; rows are not grouped by key subsets.
    """
    # Build a folder tree first so merge never jumps straight to root ".".
    merge_tree = {"items": [], "children": {}}

    for scan_item in scan_result:
        normalized_path = os.path.normpath(scan_item.source_name_or_path).replace("\\", "/")
        path_parts = [part for part in normalized_path.split("/") if part and part != "."]
        current_node = merge_tree

        for folder_name in path_parts[:-1]:
            current_node = current_node["children"].setdefault(folder_name, {"items": [], "children": {}})
        current_node["items"].append(scan_item)

    def merge_node(merge_node_item: dict, merge_path: str = "", depth: int = 0) -> tuple:
        child_finalized = []
        child_unfinalized = []
        for folder_name, child_node in merge_node_item["children"].items():
            child_path = f"{merge_path}/{folder_name}" if merge_path else folder_name
            fin, unfin = merge_node(child_node, child_path, depth + 1)
            child_finalized.extend(fin)
            child_unfinalized.extend(unfin)

        local_excluded = [item for item in merge_node_item["items"] if item.exclude]
        local_eligible = [item for item in merge_node_item["items"] if not item.exclude]

        all_eligible_candidates = list(local_eligible)
        for _, g_items in child_unfinalized:
            all_eligible_candidates.extend(g_items)

        # We can merge under the current node if depth > 0 and we are combining multiple sources:
        # e.g., local files + child groups, or multiple child groups, or multiple local files.
        can_merge_here = (depth > 0) and (len(local_eligible) + len(child_unfinalized) > 1)

        if can_merge_here:
            if _can_merge_folder(all_eligible_candidates):
                new_finalized = child_finalized + local_excluded
                new_unfinalized = [(merge_path, all_eligible_candidates)]
                return new_finalized, new_unfinalized
            else:
                # Compatibility broke at this level. We must finalize the subtrees.
                finalized_child_groups = []
                for c_path, c_items in child_unfinalized:
                    if len(c_items) > 1 and _can_merge_folder(c_items):
                        finalized_child_groups.append(_create_merged_item(c_items, c_path))
                    else:
                        finalized_child_groups.extend(c_items)

                finalized_local = []
                if len(local_eligible) > 1 and _can_merge_folder(local_eligible):
                    finalized_local.append(_create_merged_item(local_eligible, merge_path))
                else:
                    finalized_local.extend(local_eligible)

                new_finalized = child_finalized + local_excluded + finalized_child_groups + finalized_local
                new_unfinalized = []
                return new_finalized, new_unfinalized
        else:
            # We cannot merge or don't need to merge at this level (e.g. depth == 0 or only 1 source).
            # We propagate everything up as unfinalized.
            new_finalized = child_finalized + local_excluded
            new_unfinalized = list(child_unfinalized)
            if local_eligible:
                new_unfinalized.append((merge_path, local_eligible))
            return new_finalized, new_unfinalized

    fin, unfin = merge_node(merge_tree)
    finalized_results = list(fin)
    for path, items in unfin:
        if path and len(items) > 1 and _can_merge_folder(items):
            finalized_results.append(_create_merged_item(items, path))
        else:
            finalized_results.extend(items)

    return finalized_results


class _ReportSourceItem:
    """Row group the folder merge can read. One instance per source path."""

    def __init__(self, path):
        self.source_name_or_path = path
        self.exclude = False
        self.comment = ""
        self.oss_items = []
        self.download_location = []
        self.copyright = []
        self.licenses = []
        self.oss_name = ""
        self.oss_version = ""
        self._source_rows = []
        self._exclude_ready = False

    def add_row(self, row, cells):
        self._source_rows.append(list(row))
        name = _normalize_merge_text(cells["OSS Name"])
        version = _normalize_merge_text(cells["OSS Version"])
        license_text = _normalize_merge_text(cells["License"])
        download = _normalize_merge_text(cells["Download Location"])
        copyright_text = cells["Copyright Text"] or ""
        comment = cells["Comment"] or ""
        is_exclude = _normalize_merge_text(cells["Exclude"]).lower() == "exclude"

        oss = OssItem(name, version, license_text, download)
        oss.copyright = copyright_text
        oss.comment = comment
        self.oss_items.append(oss)

        if name and not self.oss_name:
            self.oss_name = name
        if version and not self.oss_version:
            self.oss_version = version
        if license_text and not self.licenses:
            self.licenses = [part.strip() for part in license_text.split(",") if part.strip()]
        if download and download not in self.download_location:
            self.download_location.append(download)
        for line in str(copyright_text).splitlines():
            line = line.strip()
            if line and line not in self.copyright:
                self.copyright.append(line)
        if comment:
            parts = [part.strip() for part in str(comment).split(COMMENT_DELIMITER) if part.strip()]
            existing = [part.strip() for part in self.comment.split(COMMENT_DELIMITER) if part.strip()] if self.comment else []
            for part in parts:
                if part not in existing:
                    existing.append(part)
            self.comment = COMMENT_DELIMITER.join(existing)
        if not self._exclude_ready:
            self.exclude = is_exclude
            self._exclude_ready = True
        else:
            self.exclude = self.exclude and is_exclude

    def set_oss_item(self):
        self.oss_items = []
        copyrights = "\n".join(self.copyright) if isinstance(self.copyright, list) else (self.copyright or "")
        downloads = self.download_location or [""]
        for url in downloads:
            item = OssItem(self.oss_name, self.oss_version, self.licenses, url)
            item.copyright = copyrights
            item.comment = self.comment
            self.oss_items.append(item)


def _sheet_cell(row, index, name):
    position = index.get(name)
    if position is None or position >= len(row) or row[position] is None:
        return ""
    return str(row[position])


def _items_from_sheet(header, data_rows):
    index = {name: position for position, name in enumerate(header)}
    grouped = {}
    order = []
    for row in data_rows:
        cells = {name: _sheet_cell(row, index, name) for name in _SRC_COLUMNS}
        path = _normalize_merge_text(str(cells["Source Path"]))
        key = path or id(row)
        if key not in grouped:
            grouped[key] = _ReportSourceItem(path)
            order.append(key)
        grouped[key].add_row(row, cells)
    return [grouped[key] for key in order]


def _merged_sheet_row(header, item, oss):
    values = {
        "Source Path": item.source_name_or_path,
        "OSS Name": oss.name,
        "OSS Version": oss.version,
        "License": ",".join(oss.license),
        "Download Location": oss.download_location or "",
        "Homepage": "",
        "Copyright Text": oss.copyright or "",
        "Exclude": "Exclude" if item.exclude else "",
        "Comment": oss.comment or "",
    }
    return [None if name == "ID" else values.get(name) for name in header]


def _rows_from_items(header, items):
    rows = []
    for item in items:
        if item._source_rows:
            rows.extend(item._source_rows)
        else:
            rows.extend(_merged_sheet_row(header, item, oss) for oss in item.oss_items)
    return rows


def _column(header, row, name):
    if name not in header:
        return ""
    position = header.index(name)
    if position >= len(row) or row[position] is None:
        return ""
    return str(row[position])


def _sort_sheet_rows(header, rows):
    return sorted(rows, key=lambda row: (
        _column(header, row, "Exclude"),
        _column(header, row, "Source Path") == "",
        _column(header, row, "Source Path"),
    ))


def _renumber_ids(header, rows):
    if "ID" not in header:
        return rows
    position = header.index("ID")
    for number, row in enumerate(rows, start=1):
        if position < len(row):
            row[position] = number
    return rows


def _merge_src_worksheet(worksheet):
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return
    header = list(rows[0])
    if "Source Path" not in header:
        return
    items = _items_from_sheet(header, rows[1:])
    merged_rows = _renumber_ids(header, _sort_sheet_rows(header, _rows_from_items(header, merge_results_by_folder(items))))
    if worksheet.max_row > 1:
        worksheet.delete_rows(2, worksheet.max_row - 1)
    for row in merged_rows:
        worksheet.append(row)


def merge_report_xlsx(path):
    """
    Return a copy of a FOSSLight report with SRC_FL_Source merged by folder.

    The file at ``path`` is left unchanged. Other sheets are copied as they are.
    """
    source = str(path)
    try:
        workbook = openpyxl.load_workbook(source)
    except Exception:
        with open(source, "rb") as report:
            buffer = io.BytesIO(report.read())
        buffer.seek(0)
        return buffer

    try:
        if SRC_SHEET_NAME in workbook.sheetnames:
            _merge_src_worksheet(workbook[SRC_SHEET_NAME])
        buffer = io.BytesIO()
        workbook.save(buffer)
    finally:
        workbook.close()
    buffer.seek(0)
    return buffer
